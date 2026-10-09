from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from app import (
    build_combined_summary_rows,
    build_combined_yearly_cumulative_rows,
    build_summary_combine_sheet,
    combine_reports_in_order,
    dataframe_to_excel_bytes,
    prepare_indicator_raw_dataframe,
)


def test_combine_reports_in_order_keeps_rows_in_input_order() -> None:
    df1 = pd.DataFrame(
        [
            {
                "Period": "Q1",
                "Organization": "PRF",
                "Project Name": "Camp Immunization",
                "indicator": "Vaccination",
                "S1 Male": 1,
                "S1 Female": 2,
            }
        ]
    )
    df2 = pd.DataFrame(
        [
            {
                "Period": "Q2",
                "Organization": "PRF",
                "Project Name": "Camp Immunization",
                "indicator": "Vaccination",
                "S1 Male": 3,
                "S1 Female": 4,
            }
        ]
    )

    combined = combine_reports_in_order(df1, df2)

    assert len(combined) == 2
    assert combined["Period"].tolist() == ["Q1", "Q2"]
    assert combined["indicator"].tolist() == ["Vaccination", "Vaccination"]


def test_build_summary_combine_sheet_contains_combined_rows() -> None:
    summary_df = pd.DataFrame(
        [
            {
                "Period": "Q1",
                "Organization": "PRF",
                "Project Name": "Camp Immunization",
                "indicator": "Vaccination",
                "S1 Male": 1,
                "S1 Female": 2,
            },
            {
                "Period": "Q1",
                "Organization": "PRF",
                "Project Name": "Camp Immunization",
                "indicator": "Vaccination",
                "S1 Male": 3,
                "S1 Female": 4,
            },
        ]
    )

    combined_sheet = build_summary_combine_sheet(summary_df)

    assert len(combined_sheet) == 1
    assert combined_sheet.iloc[0]["S1 Male"] == 4
    assert combined_sheet.iloc[0]["S1 Female"] == 6


def test_dataframe_to_excel_bytes_names_combined_indicator_sheet_indicators() -> None:
    empty_df = pd.DataFrame()

    workbook = load_workbook(
        BytesIO(
            dataframe_to_excel_bytes(
                empty_df, empty_df, empty_df, empty_df, empty_df, empty_df
            )
        ),
        read_only=True,
    )

    assert workbook.sheetnames == [
        "Indicator Semester Achievement",
        "Age_semester",
        "indicators",
        "Summary",
        "yearly_cumulative",
        "Summary_combine",
    ]


def test_build_combined_summary_rows_appends_rows_from_both_summary_sheets() -> None:
    source1 = BytesIO()
    source2 = BytesIO()
    with pd.ExcelWriter(source1, engine="openpyxl") as writer:
        pd.DataFrame(
            {
                "Period": ["Q1"],
                "Project Name": ["Mae La"],
                "Value": [1],
            }
        ).to_excel(writer, sheet_name="Summary", index=False)
        pd.DataFrame({"Wrong source sheet": [True]}).to_excel(
            writer, sheet_name="cummu_summary", index=False
        )
    with pd.ExcelWriter(source2, engine="openpyxl") as writer:
        pd.DataFrame(
            {"Period": ["Q1"], "Project Name": ["Umpium"], "Value": [3]}
        ).to_excel(writer, sheet_name="Summary", index=False)
        pd.DataFrame({"Wrong source sheet": [True]}).to_excel(
            writer, sheet_name="yearly_cummu_summary", index=False
        )
    source1.seek(0)
    source2.seek(0)

    result = build_combined_summary_rows(source1, source2)

    assert result["Project Name"].tolist() == ["Mae La", "Umpium"]
    assert result["Value"].tolist() == [1, 3]
    assert "Wrong source sheet" not in result.columns


def test_build_combined_yearly_cumulative_rows_uses_cumulative_source_sheets() -> None:
    source1 = BytesIO()
    source2 = BytesIO()
    with pd.ExcelWriter(source1, engine="openpyxl") as writer:
        pd.DataFrame(
            {
                "Year": [2026],
                "Organization": ["Mae La"],
                "Project Name": ["Camp Immunization"],
                "Twp_MIMU": ["Mae La Township"],
                "Clinic Name": ["Clinic A"],
                "ALOD_U1": [1],
            }
        ).to_excel(writer, sheet_name="cummu_summary", index=False)
        pd.DataFrame({"Wrong source sheet": [True]}).to_excel(
            writer, sheet_name="Summary", index=False
        )
    with pd.ExcelWriter(source2, engine="openpyxl") as writer:
        pd.DataFrame(
            {
                "Year": [2026],
                "Organization": ["PRF"],
                "Project Name": ["Camp Immunization"],
                "District (EHO)": ["District A"],
                "Township_EHO": ["Township A"],
                "Twp_MIMU": ["Township A"],
                "Clinic Name": ["Clinic B"],
                "ALOD_U1": [3],
            }
        ).to_excel(writer, sheet_name="yearly_cummu_summary", index=False)
        pd.DataFrame({"Wrong source sheet": [True]}).to_excel(
            writer, sheet_name="Summary", index=False
        )
    source1.seek(0)
    source2.seek(0)

    result = build_combined_yearly_cumulative_rows(source1, source2)

    assert result["Organization"].tolist() == ["Mae La", "PRF"]
    assert result["District (EHO)"].tolist() == ["", "District A"]
    assert result["Township_EHO"].tolist() == ["", "Township A"]
    assert result["ALOD_U1"].tolist() == [1, 3]
    assert "Wrong source sheet" not in result.columns


def test_prepare_indicator_raw_dataframe_uses_reach_kk_project_name() -> None:
    source = BytesIO()
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        pd.DataFrame(
            {
                "Period": ["Q1"],
                "indicator": ["Vaccination"],
                "Value": [1],
            }
        ).to_excel(writer, sheet_name="indicator", index=False)
    source.seek(0)

    result = prepare_indicator_raw_dataframe(source)

    assert result["Project Name"].tolist() == ["REACH-KK"]
